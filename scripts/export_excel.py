#!/usr/bin/env python3
"""Export validated lead-keyword JSON to a styled Excel workbook."""

import argparse
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

LEVELS = ("priority_1", "priority_2", "priority_3")
ITEM_KEYS = ("keyword", "logic", "market_size", "trend")
NAVY = "102A43"
GOLD = "C9A227"
WHITE = "FFFFFF"
TEXT = "243B53"
MUTED = "627D98"
PALE = "F0F4F8"


def fill(color):
    return PatternFill("solid", fgColor=color)


def style_table(ws, widths):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:D{ws.max_row}"
    ws.row_dimensions[1].height = 30
    for cell in ws[1]:
        cell.fill = fill(NAVY)
        cell.font = Font(name="微软雅黑", size=11, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="D9E2EC")
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 66
        for col in range(1, 5):
            cell = ws.cell(row, col)
            cell.font = Font(name="微软雅黑", size=10, color=TEXT)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
            if row % 2 == 0:
                cell.fill = fill(PALE)
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def add_summary(wb, summary, counts):
    ws = wb.active
    ws.title = "business_summary"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:D1")
    ws["A1"] = "地图获客行业关键词分析"
    ws["A1"].fill = fill(NAVY)
    ws["A1"].font = Font(name="微软雅黑", size=20, bold=True, color=WHITE)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 54
    ws["A3"] = "business_summary"
    ws["A3"].font = Font(name="微软雅黑", size=11, bold=True, color=GOLD)
    ws.merge_cells("A4:D7")
    ws["A4"] = summary
    ws["A4"].font = Font(name="微软雅黑", size=12, color=TEXT)
    ws["A4"].alignment = Alignment(vertical="top", wrap_text=True)
    ws["A4"].fill = fill(PALE)
    ws["A9"] = "工作表"
    ws["B9"] = "关键词数量"
    for cell in ws[9][:2]:
        cell.fill = fill(NAVY)
        cell.font = Font(name="微软雅黑", size=11, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center")
    for row, level in enumerate(LEVELS, 10):
        ws.cell(row, 1, level)
        ws.cell(row, 2, counts[level])
        for col in (1, 2):
            ws.cell(row, col).font = Font(name="微软雅黑", size=11, color=TEXT)
            ws.cell(row, col).alignment = Alignment(horizontal="center")
    ws["A15"] = "趋势计分参考"
    ws["A15"].font = Font(name="微软雅黑", size=11, bold=True, color=GOLD)
    ws["A16"] = "上涨 +2｜温和上涨 +1｜结构性分化 0｜下行 -2（不计入关键词字段）"
    ws.merge_cells("A16:D16")
    ws["A16"].font = Font(name="微软雅黑", size=10, color=MUTED)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 26
    ws.column_dimensions["D"].width = 26


def add_sources(wb, sources):
    ws = wb.create_sheet("资料来源")
    ws.sheet_view.showGridLines = False
    headers = ("source_name", "source_url", "note")
    ws.append(headers)
    for source in sources:
        ws.append([source.get(key, "") for key in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{ws.max_row}"
    for cell in ws[1]:
        cell.fill = fill(NAVY)
        cell.font = Font(name="微软雅黑", size=11, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center")
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 54
        for col in range(1, 4):
            ws.cell(row, col).font = Font(name="微软雅黑", size=10, color=TEXT)
            ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
    for col, width in enumerate((34, 72, 48), 1):
        ws.column_dimensions[get_column_letter(col)].width = width


def export(data, output_path, sources=None):
    wb = Workbook()
    counts = {level: len(data[level]) for level in LEVELS}
    add_summary(wb, data["business_summary"], counts)
    for level in LEVELS:
        ws = wb.create_sheet(level)
        ws.append(ITEM_KEYS)
        for item in data[level]:
            ws.append([item[key] for key in ITEM_KEYS])
        style_table(ws, (20, 66, 42, 42))
    if sources:
        add_sources(wb, sources)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    check = load_workbook(output_path, read_only=True, data_only=False)
    assert check["business_summary"]["A4"].value == data["business_summary"]
    for level in LEVELS:
        assert check[level].max_row == len(data[level]) + 1
        assert tuple(cell.value for cell in check[level][1]) == ITEM_KEYS
    check.close()


def main():
    parser = argparse.ArgumentParser(description="将地图获客关键词 JSON 导出为 Excel")
    parser.add_argument("input_json", type=Path)
    parser.add_argument("output_xlsx", type=Path)
    parser.add_argument("--sources", type=Path, help="可选：source_name/source_url/note 数组 JSON")
    args = parser.parse_args()
    data = json.loads(args.input_json.read_text(encoding="utf-8-sig"))
    sources = None
    if args.sources:
        sources = json.loads(args.sources.read_text(encoding="utf-8-sig"))
        if not isinstance(sources, list):
            raise ValueError("sources 文件必须是 JSON 数组")
    export(data, args.output_xlsx, sources)
    print(f"Excel 已生成：{args.output_xlsx}")


if __name__ == "__main__":
    main()
